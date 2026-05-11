#!/usr/bin/env python3
"""
Claude Skills QA Framework – Eval Runner v1.1 by Kentobayashi
──────────────────────────────────────────────────────────────
Automatically tests Claude Skills via the Anthropic API.
Tests trigger accuracy, output quality, false positives, and consistency.

Requirements:
    pip install anthropic openpyxl pyyaml

Setup:
    1. Set your Anthropic API key as environment variable:
       Windows CMD:        set ANTHROPIC_API_KEY=sk-ant-...
       Windows PowerShell: $env:ANTHROPIC_API_KEY="sk-ant-..."
       macOS/Linux:        export ANTHROPIC_API_KEY=sk-ant-...

    2. Set REPO_PATH to your local skill library folder (see CONFIGURATION below)

    3. Add your skills to skill_eval_data.py

Usage:
    python skill_eval_runner.py                    # Test all skills
    python skill_eval_runner.py --skill my-skill   # Test one skill
    python skill_eval_runner.py --category Core    # Test by category
    python skill_eval_runner.py --dry-run          # No API calls, check structure only
"""

import os, sys, json, time, re, argparse
from pathlib import Path
from datetime import datetime

# ═══════════════════════════════════════════════════════════════════
# CONFIGURATION – adjust these values for your setup
# ═══════════════════════════════════════════════════════════════════

REPO_PATH   = Path("./skills")          # Path to your local skill library
                                        # Skills should be at: REPO_PATH / skill["folder"] / "SKILL.md"
API_KEY     = os.environ.get("ANTHROPIC_API_KEY", "")

MODEL_FAST  = "claude-haiku-4-5-20251001"   # Trigger detection + Judge (cost-efficient)
MODEL_MAIN  = "claude-sonnet-4-6"           # Output generation (quality)

DELAY_SEC   = 0.8     # Pause between API calls (rate limiting)
MAX_RETRIES = 3       # Retries on 429/500 errors
MAX_TOKENS_OUTPUT = 8000
MAX_TOKENS_FAST   = 150

OUTPUT_FILE = f"skill-eval-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"

# ═══════════════════════════════════════════════════════════════════
# DEPENDENCY CHECK
# ═══════════════════════════════════════════════════════════════════

try:
    import anthropic
except ImportError:
    sys.exit("❌ Missing: pip install anthropic")

try:
    import yaml
except ImportError:
    sys.exit("❌ Missing: pip install pyyaml")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import ColorScaleRule
except ImportError:
    sys.exit("❌ Missing: pip install openpyxl")

from skill_eval_data import SKILLS

# ═══════════════════════════════════════════════════════════════════
# HELPER FUNCTIONS
# ═══════════════════════════════════════════════════════════════════

def log(msg, level="INFO"):
    prefix = {"INFO": "  ·", "OK": "  ✅", "WARN": "  ⚠️", "ERR": "  ❌", "HEAD": "\n══"}
    print(f"{prefix.get(level, '  ')} {msg}")

def read_skill_md(skill: dict) -> str | None:
    """Reads SKILL.md from the local skill library."""
    path = REPO_PATH / skill["folder"] / "SKILL.md"
    if not path.exists():
        log(f"SKILL.md not found: {path}", "WARN")
        return None
    return path.read_text(encoding="utf-8")

def extract_description(skill_content: str) -> str:
    """Extracts the description field from YAML frontmatter."""
    match = re.match(r"^---\n(.*?)\n---", skill_content, re.DOTALL)
    if not match:
        return skill_content[:500]
    try:
        fm = yaml.safe_load(match.group(1))
        return str(fm.get("description", ""))[:800]
    except Exception:
        return skill_content[:500]

def api_call(client, model: str, system: str, user: str, max_tokens: int) -> str:
    """API call with retry logic."""
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = client.messages.create(
                model=model,
                max_tokens=max_tokens,
                system=system,
                messages=[{"role": "user", "content": user}],
            )
            return response.content[0].text.strip()
        except anthropic.RateLimitError:
            wait = 10 * attempt
            log(f"Rate limit – waiting {wait}s (attempt {attempt}/{MAX_RETRIES})", "WARN")
            time.sleep(wait)
        except anthropic.APIStatusError as e:
            log(f"API error {e.status_code}: {e.message} (attempt {attempt})", "WARN")
            time.sleep(5 * attempt)
        except Exception as e:
            log(f"Unexpected error: {e} (attempt {attempt})", "WARN")
            time.sleep(3)
    return ""

# ═══════════════════════════════════════════════════════════════════
# EVAL STAGES
# ═══════════════════════════════════════════════════════════════════

def test_trigger(client, description: str, prompt: str) -> bool:
    """Stage T: Should the skill activate on this input? → YES / NO"""
    system = (
        "You decide whether a Claude Skill should activate based on its description. "
        "Reply with exactly one word: YES or NO."
    )
    user = (
        f"Skill description:\n{description}\n\n"
        f"User message:\n{prompt}\n\n"
        "Should this skill activate? Reply YES or NO only."
    )
    result = api_call(client, MODEL_FAST, system, user, MAX_TOKENS_FAST)
    time.sleep(DELAY_SEC)
    return "YES" in result.upper()

def generate_output(client, skill_content: str, prompt: str) -> str:
    """Stage O: Generate skill output for the trigger prompt."""
    result = api_call(client, MODEL_MAIN, skill_content, prompt, MAX_TOKENS_OUTPUT)
    time.sleep(DELAY_SEC)
    return result

def judge_output(client, output: str, criteria: str) -> tuple[int, str]:
    """Stage O/K: Evaluate output quality – returns (score 0-5, reason)."""
    system = (
        "You are an expert evaluator for Claude AI skill outputs. "
        "Score the output on a 0-5 scale based on how well it meets the given criteria. "
        "0 = completely misses criteria, 5 = perfectly meets all criteria. "
        'Respond ONLY with valid JSON: {"score": <integer 0-5>, "reason": "<max 20 words>"}'
    )
    user = (
        f"Output to evaluate (first 1500 chars):\n{output[:1500]}\n\n"
        f"Quality criteria:\n{criteria}\n\n"
        "Score 0-5 and give a brief reason."
    )
    raw = api_call(client, MODEL_FAST, system, user, MAX_TOKENS_FAST)
    time.sleep(DELAY_SEC)
    try:
        json_match = re.search(r'\{.*?\}', raw, re.DOTALL)
        if json_match:
            data = json.loads(json_match.group())
            score = max(0, min(5, int(data.get("score", 0))))
            return score, str(data.get("reason", ""))[:120]
    except Exception:
        pass
    return 0, f"Parse error: {raw[:60]}"

# ═══════════════════════════════════════════════════════════════════
# SKILL EVALUATION
# ═══════════════════════════════════════════════════════════════════

def eval_skill(client, skill: dict, dry_run: bool = False) -> dict:
    """
    Runs all tests for one skill.
    Returns dict with T, O, FP, K, total, status, notes.

    Scoring:
        T  (Trigger quality):    3/3 hits → 5pts, 2/3 → 3pts, 1/3 → 1pt,  0/3 → 0pts
        FP (False positive):     2/2 OK   → 5pts, 1/2 → 2pts, 0/2 → 0pts
        O  (Output quality):     0-5pts from LLM judge
        K  (Consistency):        5 - abs(O_run1 - O_run2) * 2
        Total: T + O + FP + K (max 20)

    Thresholds:
        ≥16/20: ✅ Released
        11-15:  ⚠️  Fix needed
        ≤10:    🔴 Rebuild
    """
    name = skill["name"]
    log(f"Testing: {name} [{skill['category']}]", "HEAD")

    skill_content = read_skill_md(skill)
    if not skill_content:
        return {"name": name, "category": skill["category"],
                "T": 0, "O": 0, "FP": 0, "K": 0, "total": 0,
                "status": "🔴 Rebuild", "notes": "SKILL.md not found"}

    description = extract_description(skill_content)

    if dry_run:
        log("Dry-run – skipping API calls", "WARN")
        return {"name": name, "category": skill["category"],
                "T": "-", "O": "-", "FP": "-", "K": "-", "total": "-",
                "status": "⏭️ Dry-Run", "notes": "No API call"}

    # ── T: TRIGGER QUALITY ───────────────────────────────────────
    log("T: Positive trigger tests (3x)...")
    t_hits = 0
    for i, prompt in enumerate(skill["trigger_pos"]):
        hit = test_trigger(client, description, prompt)
        log(f"  Trigger {i+1}: {'✅ YES' if hit else '❌ NO'} – {prompt[:60]}...")
        if hit:
            t_hits += 1

    T = {3: 5, 2: 3, 1: 1, 0: 0}[t_hits]
    log(f"  → T = {T}/5 ({t_hits}/3 triggered)")

    # ── FP: FALSE POSITIVES ──────────────────────────────────────
    log("FP: Negative trigger tests (2x)...")
    fp_correct = 0
    for i, prompt in enumerate(skill["trigger_neg"]):
        hit = test_trigger(client, description, prompt)
        correct = not hit
        log(f"  Negative {i+1}: {'✅ no trigger (correct)' if correct else '❌ false positive!'} – {prompt[:60]}...")
        if correct:
            fp_correct += 1

    FP = {2: 5, 1: 2, 0: 0}[fp_correct]
    log(f"  → FP = {FP}/5 ({fp_correct}/2 correct)")

    # ── O: OUTPUT QUALITY ────────────────────────────────────────
    log("O: Output generation + evaluation...")
    output1 = generate_output(client, skill_content, skill["trigger_pos"][0])
    O, o_reason = judge_output(client, output1, skill["output_criteria"])
    log(f"  → O = {O}/5 | {o_reason}")

    # ── K: CONSISTENCY ───────────────────────────────────────────
    log("K: Consistency test (same trigger, 2nd run)...")
    output2 = generate_output(client, skill_content, skill["trigger_pos"][0])
    O2, _ = judge_output(client, output2, skill["output_criteria"])
    K = max(0, 5 - abs(O - O2) * 2)
    log(f"  → K = {K}/5 (Run1={O}, Run2={O2})")

    # ── TOTAL ────────────────────────────────────────────────────
    total = T + O + FP + K
    if total >= 16:
        status = "✅ OK"
    elif total >= 11:
        status = "⚠️ Fix"
    else:
        status = "🔴 Rebuild"

    notes = f"T-Hits:{t_hits}/3 | FP-OK:{fp_correct}/2 | O-Run2:{O2} | {o_reason[:60]}"
    log(f"  TOTAL: {total}/20 → {status}")

    return {
        "name": name,
        "category": skill["category"],
        "T": T, "O": O, "FP": FP, "K": K,
        "total": total, "status": status, "notes": notes,
    }

# ═══════════════════════════════════════════════════════════════════
# XLSX OUTPUT
# ═══════════════════════════════════════════════════════════════════

def save_results(results: list[dict], filepath: str):
    """Writes results to a formatted XLSX file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Eval Results"
    ws.sheet_view.showGridLines = False

    C_DARK   = "1A1A2E"
    C_ACCENT = "E94560"
    C_OFFWH  = "F8F9FA"
    C_SCORE  = "EBF5FB"
    C_GOLD   = "FFF3CD"

    def fill(h): return PatternFill("solid", fgColor=h)
    def f(bold=False, size=10, color="000000"):
        return Font(name="Arial", bold=bold, size=size, color=color)
    def ctr(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
    def lft(): return Alignment(horizontal="left",   vertical="center", wrap_text=True)
    def brd():
        s = Side(style="thin", color="CCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)

    for col, w in [(1,5),(2,26),(3,10),(4,8),(5,8),(6,8),(7,8),(8,10),(9,14),(10,40)]:
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 40
    ws.merge_cells("A1:J1")
    c = ws["A1"]
    c.value = f"📊 Claude Skills QA Framework – {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    c.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    c.fill = fill(C_DARK)
    c.alignment = ctr()

    ws.row_dimensions[2].height = 22
    ok_count  = sum(1 for r in results if "✅" in str(r.get("status","")))
    fix_count = sum(1 for r in results if "⚠️" in str(r.get("status","")))
    neu_count = sum(1 for r in results if "🔴" in str(r.get("status","")))
    valid = [r for r in results if isinstance(r.get("total"), int)]
    avg = round(sum(r["total"] for r in valid) / len(valid), 1) if valid else "-"

    ws.merge_cells("A2:J2")
    c = ws["A2"]
    c.value = (f"Total: {len(results)} Skills | "
               f"✅ Released: {ok_count} | ⚠️ Fix needed: {fix_count} | "
               f"🔴 Rebuild: {neu_count} | Ø Score: {avg}/20")
    c.font = Font(name="Arial", size=10, color="FFFFFF")
    c.fill = fill("0F3460")
    c.alignment = ctr()

    headers = ["#", "Skill Name", "Category", "T /5", "O /5", "FP /5", "K /5", "Total /20", "Status", "Notes"]
    ws.row_dimensions[4].height = 24
    for j, h in enumerate(headers):
        c = ws.cell(row=4, column=j+1, value=h)
        c.font = f(bold=True, size=9, color="FFFFFF")
        c.fill = fill(C_DARK)
        c.alignment = ctr()
        c.border = brd()

    ws.freeze_panes = "A5"

    for i, r in enumerate(results):
        row = 5 + i
        ws.row_dimensions[row].height = 28
        bg = C_OFFWH if i % 2 == 0 else "FFFFFF"

        values = [
            i + 1, r["name"], r["category"],
            r.get("T","-"), r.get("O","-"), r.get("FP","-"), r.get("K","-"),
            r.get("total","-"), r.get("status","-"), r.get("notes",""),
        ]
        for j, val in enumerate(values):
            c = ws.cell(row=row, column=j+1, value=val)
            c.border = brd()
            c.alignment = lft() if j in [1, 9] else ctr()

            if j == 0:
                c.font = f(bold=True, size=10, color="FFFFFF")
                c.fill = fill(C_ACCENT)
            elif j == 1:
                c.font = f(bold=True, size=10)
                c.fill = fill(bg)
            elif j in [3,4,5,6]:
                c.font = f(bold=True, size=12)
                c.fill = fill(C_SCORE)
            elif j == 7:
                c.font = f(bold=True, size=13)
                c.fill = fill(C_GOLD)
            elif j == 8:
                c.font = f(bold=True, size=11)
                c.fill = fill(C_SCORE)
            else:
                c.font = f(size=8)
                c.fill = fill(bg)

    last = 4 + len(results)
    ws.conditional_formatting.add(
        f"H5:H{last}",
        ColorScaleRule(
            start_type="num", start_value=0,  start_color="E74C3C",
            mid_type="num",   mid_value=12,   mid_color="F39C12",
            end_type="num",   end_value=20,   end_color="27AE60",
        ),
    )

    wb.save(filepath)
    log(f"Results saved: {filepath}", "OK")

# ═══════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(description="Claude Skills QA Framework – Eval Runner")
    parser.add_argument("--skill",    help="Test one skill by name (e.g. my-skill)")
    parser.add_argument("--category", help="Test by category (e.g. Core or Light)")
    parser.add_argument("--dry-run",  action="store_true", help="Check structure without API calls")
    args = parser.parse_args()

    print("\n" + "═"*60)
    print("  Claude Skills QA Framework – Eval Runner v1.1")
    print("  by Kentobayashi")
    print("═"*60)

    if not args.dry_run:
        if not API_KEY:
            sys.exit("❌ ANTHROPIC_API_KEY not set. Please set it as an environment variable.")
        client = anthropic.Anthropic(api_key=API_KEY)
        log("Connected to Anthropic API", "OK")
        log(f"Trigger/Judge model: {MODEL_FAST}")
        log(f"Output model:        {MODEL_MAIN}")
    else:
        client = None
        log("DRY-RUN mode – no API calls", "WARN")

    skills_to_test = SKILLS
    if args.skill:
        skills_to_test = [s for s in SKILLS if s["name"] == args.skill]
        if not skills_to_test:
            sys.exit(f"❌ Skill '{args.skill}' not found in skill_eval_data.py")
    if args.category:
        skills_to_test = [s for s in skills_to_test if s["category"].lower() == args.category.lower()]

    log(f"Skills to test: {len(skills_to_test)}")
    if not args.dry_run:
        est_calls = len(skills_to_test) * 9
        est_cost  = round(len(skills_to_test) * 0.18, 2)
        est_time  = round(len(skills_to_test) * 45 / 60, 1)
        log(f"Estimated API calls: ~{est_calls}")
        log(f"Estimated cost:      ~${est_cost}")
        log(f"Estimated time:      ~{est_time} minutes")
        print()
        confirm = input("  Start? (y/n): ").strip().lower()
        if confirm != "y":
            sys.exit("Cancelled.")

    results = []
    start = time.time()

    for i, skill in enumerate(skills_to_test):
        print(f"\n[{i+1}/{len(skills_to_test)}]", end="")
        result = eval_skill(client, skill, dry_run=args.dry_run)
        results.append(result)

    save_results(results, OUTPUT_FILE)

    elapsed = round(time.time() - start, 1)
    print("\n" + "═"*60)
    print(f"  ✅ Eval completed in {elapsed}s")
    valid = [r for r in results if isinstance(r.get("total"), int)]
    if valid:
        avg = round(sum(r["total"] for r in valid) / len(valid), 1)
        ok  = sum(1 for r in valid if "✅" in str(r["status"]))
        fix = sum(1 for r in valid if "⚠️" in str(r["status"]))
        neu = sum(1 for r in valid if "🔴" in str(r["status"]))
        print(f"  Ø Score:       {avg}/20")
        print(f"  ✅ Released:   {ok} Skills")
        print(f"  ⚠️  Fix needed: {fix} Skills")
        print(f"  🔴 Rebuild:    {neu} Skills")
    print(f"  📄 Results: {OUTPUT_FILE}")
    print("═"*60)

if __name__ == "__main__":
    main()
